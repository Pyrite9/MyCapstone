# 대본 v1.2 빌더: 이전 대본(v2.5 모듈형)의 문단 서식을 복사해 새 대본 docx와 대조표 docx를 만든다
# 사용: python3 build_script_v12.py 대본템플릿.docx 대조표템플릿.docx 대본출력.docx 대조표출력.docx
import copy, re, sys
from docx import Document
sys.path.insert(0, '/home/claude/w10_3')
from script_data_v12 import META, TITLE, SUBTITLE, SECTIONS, QA

TPL, CMP_TPL, OUT, CMP_OUT = sys.argv[1:5]
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
XS = '{http://www.w3.org/XML/1998/namespace}space'
G_LINES = open('/home/claude/w10_3/기능도기술서_원고_v1.2.txt', encoding='utf-8').read().split('\n')
O_LINES = open('/mnt/user-data/uploads/캡스톤/보고서/개요서_원고_v6.5.txt', encoding='utf-8').read().split('\n')


def set_text(r, t):
    for tt in r.findall(W + 't'):
        r.remove(tt)
    te = r.makeelement(W + 't', {}); te.text = t; te.set(XS, 'preserve')
    r.append(te)


# ---------- 대본 ----------
d = Document(TPL)
ps = d.paragraphs
S = {'blank': ps[0], 'title': ps[1], 'sub': ps[2], 'meta': ps[4], 'pb': ps[9],
     'sec': ps[10], 'mod': ps[11], 'body': ps[12], 'q': ps[97], 'a': ps[98]}
S = {k: copy.deepcopy(v._p) for k, v in S.items()}
body = d.element.body
sectPr = body.find(W + 'sectPr')
for el in list(body):
    if el is not sectPr:
        body.remove(el)


def add(kind, texts=None):
    p = copy.deepcopy(S[kind])
    runs = p.findall(W + 'r')
    if texts is not None:
        for r, t in zip(runs, texts):
            set_text(r, t)
        for r in runs[len(texts):]:
            p.remove(r)
    sectPr.addprevious(p)
    return p


def add_body(text, kind='body'):
    p = copy.deepcopy(S[kind])
    runs = p.findall(W + 'r')
    base = runs[0]
    for r in runs:
        p.remove(r)
    for i, seg in enumerate(re.split(r'\*\*', text)):
        if not seg:
            continue
        r = copy.deepcopy(base)
        rpr = r.find(W + 'rPr')
        b = rpr.find(W + 'b') if rpr is not None else None
        if i % 2 == 1:
            if b is None:
                rpr.insert(0, rpr.makeelement(W + 'b', {}))
        elif b is not None:
            rpr.remove(b)
        set_text(r, seg); p.append(r)
    sectPr.addprevious(p)


add('blank')
add('title', [TITLE]); add('sub', [SUBTITLE]); add('blank')
for k, v in META:
    add('meta', [k, '   ' + v])
add('pb')
for sec, mods in SECTIONS:
    add('sec', [sec])
    for head, name, paras, src in mods:
        add('mod', [head + '  ', name])
        for t in paras:
            add_body(t)
add('pb')
add('sec', ['예상 질문과 답변'])
for q, answers, src in QA:
    add('q', ['Q. ' + q])
    for a in answers:
        add_body(a, 'a')
d.save(OUT)
print('saved', OUT)

# ---------- 대조표 ----------
c = Document(CMP_TPL)
cbody = c.element.body
els = list(cbody)
csect = cbody.find(W + 'sectPr')
T = {'title': els[0], 'note': els[1], 'sec': els[2], 'mod': els[3], 'tbl': els[4]}
T = {k: copy.deepcopy(v) for k, v in T.items()}
for el in list(cbody):
    if el is not csect:
        cbody.remove(el)


def cadd(kind, text):
    p = copy.deepcopy(T[kind])
    runs = p.findall(W + 'r')
    set_text(runs[0], text)
    for r in runs[1:]:
        p.remove(r)
    csect.addprevious(p)


def clean(line):
    tag, _, rest = line.partition(' ')
    if tag in ('O', 'D', 'N', 'H', 'S'):
        return rest
    if tag == 'TABLE':
        return rest
    if tag == 'BOX':
        return '[요약 상자] ' + re.sub(r'\s+', ' ', rest)
    if tag == 'TITLE':
        return '[제목] ' + rest
    if tag == 'SUB':
        return '[부제] ' + rest
    if tag == 'COVER':
        return '[표지] ' + rest
    if tag == 'FIG':
        f = rest.split('|')
        return f[1] + ' (그림)'
    if line in ('END', 'PAGE', 'TOC') or line.startswith('T1 ') or line.startswith('T2 '):
        return None
    return line


def src_lines(src):
    out = []
    for item in src:
        if item[0] == 'G':
            for i in range(item[1], item[2] + 1):
                t = clean(G_LINES[i - 1].strip())
                if t:
                    out.append(t)
        elif item[0] == 'O':
            for i in range(item[1], item[2] + 1):
                t = clean(O_LINES[i - 1].strip())
                if t:
                    out.append('[개요서 %d행] ' % i + t)
        else:
            out.append('[결정] ' + item[1])
    return out


def src_label(src):
    g = [(a, b) for k, a, b in [x for x in src if x[0] in 'GO'] if k == 'G']
    o = [(a, b) for k, a, b in [x for x in src if x[0] in 'GO'] if k == 'O']
    parts = []
    def fmt(rs):
        return ', '.join(('%d' % a if a == b else '%d~%d' % (a, b)) for a, b in rs)
    if g:
        parts.append('기술서 원고 ' + fmt(g) + '행')
    if o:
        parts.append('개요서 원고 ' + fmt(o) + '행')
    if any(x[0] == 'D' for x in src):
        parts.append('결정')
    return ', '.join(parts)


def ctable(left, right):
    t = copy.deepcopy(T['tbl'])
    tcs = t.findall('.//' + W + 'tc')
    for tc, lines in ((tcs[0], left), (tcs[1], right)):
        ps_ = tc.findall(W + 'p')
        sample = copy.deepcopy(ps_[0])
        for p in ps_:
            tc.remove(p)
        for line in (lines or ['(해당 없음)']):
            p = copy.deepcopy(sample)
            runs = p.findall(W + 'r')
            base = runs[0]
            for r in runs:
                p.remove(r)
            for i, seg in enumerate(re.split(r'\*\*', line)):
                if not seg:
                    continue
                r = copy.deepcopy(base)
                rpr = r.find(W + 'rPr')
                b = rpr.find(W + 'b')
                if i % 2 == 1 and b is None:
                    rpr.insert(0, rpr.makeelement(W + 'b', {}))
                elif i % 2 == 0 and b is not None:
                    rpr.remove(b)
                set_text(r, seg); p.append(r)
            tc.append(p)
    csect.addprevious(t)


cadd('title', '발자국 발표 대본 v1.2 대조표')
cadd('note', '왼쪽은 발표 자료(기능도기술서 원고 v1.2, 필요한 곳은 개요서 원고 v6.5와 결정 사항), 오른쪽은 대본. 작성일 2026.10.15')
for sec, mods in SECTIONS:
    cadd('sec', sec)
    for head, name, paras, src in mods:
        cadd('mod', '%s · %s  (%s)' % (head, name, src_label(src)))
        ctable(src_lines(src), paras)
cadd('sec', '예상 질문과 답변')
for q, answers, src in QA:
    lab = src_label(src)
    cadd('mod', 'Q. ' + q + ('  (' + lab + ')' if lab else ''))
    ctable(src_lines(src), answers)
c.save(CMP_OUT)
print('saved', CMP_OUT)
